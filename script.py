from docx import Document
import xlrd

DOCX_TEMPLATE = "template.docx"
DOCX_RESULT = "resultat.docx"
XLS_CODES = 'codes.xls'
CHARACTERS_TO_REPLACE = "aaaa-aaaa"
NB_CODES = 300

def replace_first(previous,new,text):
    """
    remplace la premiere occurence rencontrée
    """
    return text.replace(previous,new,1)

def sheet_to_list(sheet):
    result = []
    for i in range(NB_CODES):
        result.append(sheet.cell(i, 0).value)
    return result

def code_generator(codes):
    """
    renvoi le code depuis la liste de codes
    """
    count = 0
    while True:
        print(str(codes[count]))
        yield str(codes[count])
        count += 1


if __name__ == '__main__':
    # Ouverture du fichier xls contenant les codes dans la première colonne
    workbook = xlrd.open_workbook(XLS_CODES)
    # Ouverture du premier onglet excel
    worksheet = workbook.sheet_by_index(0)
    # Creation de la liste de codes récupérés d'excel
    codes = sheet_to_list(worksheet)
    # Instanciation du generateur de code depuis la liste
    generator = code_generator(codes)
    # Chargement du template via librairie docx
    document = Document(DOCX_TEMPLATE)

    count = int(NB_CODES)
    for paragraph in document.paragraphs:
        if count == 0:
            break

        # Compte le nombre de fois qu'il faut effectuer le remplacement
        # dans le paragraphe en cours

        is_found = paragraph.text.count(CHARACTERS_TO_REPLACE)

        while is_found > 0:
            if count == 0:
                break
            print("trouvé !", count)
            paragraph.text = replace_first(CHARACTERS_TO_REPLACE,
                                                next(generator),
                                                paragraph.text)
            count -= 1
            is_found -= 1





    document.save(DOCX_RESULT)
